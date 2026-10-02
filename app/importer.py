"""Read-only imports of real ticker sheets and historical CSV evidence."""
import csv
import hashlib
import re
from datetime import date, datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from app.db import connect, init_db
from app.prices import insert_snapshot, normalize_rows, validate_symbol


def _hash(path):
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _header(value):
    text = str(value or "").strip().lower()
    text = re.sub(r"\s*\([^)]*\)", "", text)
    return text.replace(" ", "_").replace("adjusted_close", "adj_close")


def _date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return date.fromisoformat(str(value).strip()[:10]).isoformat()


def _parse_rows(headers, records):
    columns = [_header(value) for value in headers]
    if not {"date", "close", "adj_close"}.issubset(columns):
        raise ValueError("실제 가격의 date, close, adj_close 열이 없습니다.")
    prices = {}
    invalid, duplicates = 0, 0
    samples = []
    source_url, retrieved = None, None
    for line, values in enumerate(records, 2):
        if all(value is None or value == "" for value in values):
            continue
        source = dict(zip(columns, values))
        try:
            row = {"date": _date(source["date"])}
            for field in ("open", "high", "low", "close", "adj_close", "volume"):
                value = source.get(field)
                row[field] = None if value in (None, "") else value
            row = normalize_rows([row])[0]
        except (ValueError, TypeError, KeyError, OverflowError) as exc:
            invalid += 1
            if len(samples) < 3:
                samples.append(f"{line}행: {exc}")
            continue
        if row["date"] in prices:
            duplicates += 1
        prices[row["date"]] = row
        source_url = source_url or source.get("source_url")
        retrieved = retrieved or source.get("retrieved_utc")
    rows = [prices[key] for key in sorted(prices)]
    if not rows:
        raise ValueError("유효한 실제 가격 행이 없습니다.")
    return rows, invalid, duplicates, samples, source_url, retrieved


def import_file(path, file_path):
    file_path = Path(file_path)
    # Existing raw sources use `QLD raw data.xlsx` and `QLD_Yahoo_...csv`.
    if file_path.suffix.lower() not in (".xlsx", ".csv"):
        return {"file": str(file_path), "status": "skipped", "reason": "지원하지 않는 파일 형식"}
    try:
        symbol = validate_symbol(re.split(r"[ _]", file_path.stem)[0])
    except ValueError:
        return {"file": str(file_path), "status": "skipped", "reason": "종목 코드로 시작하는 가격 원본 파일이 아님"}
    digest = _hash(file_path)
    init_db(path)
    with connect(path) as con:
        prior = con.execute("SELECT snapshot_id FROM import_runs WHERE sha256=? AND symbol=?", (digest, symbol)).fetchone()
    if prior:
        return {"file": str(file_path), "symbol": symbol, "status": "skipped", "reason": "동일 내용이 이미 주입됨", "snapshot_id": prior[0]}
    if file_path.suffix.lower() == ".xlsx":
        book = load_workbook(file_path, read_only=True, data_only=True)
        try:
            if symbol not in book.sheetnames:
                return {"file": str(file_path), "symbol": symbol, "status": "skipped", "reason": "종목 코드와 같은 실제 가격 시트 없음"}
            records = iter(book[symbol].iter_rows(values_only=True))
            parsed = _parse_rows(next(records), records)
        finally:
            book.close()
    else:
        with file_path.open("r", encoding="utf-8-sig", newline="") as handle:
            records = csv.reader(handle)
            parsed = _parse_rows(next(records), records)
    rows, invalid, duplicates, samples, source_url, retrieved = parsed
    currency = "USD"
    with connect(path) as con:
        # Recheck under the same transaction in case a second importer completed meanwhile.
        prior = con.execute("SELECT snapshot_id FROM import_runs WHERE sha256=? AND symbol=?", (digest, symbol)).fetchone()
        if prior:
            return {"file": str(file_path), "symbol": symbol, "status": "skipped", "snapshot_id": prior[0]}
        identifier = insert_snapshot(con, symbol, rows, "raw", source_url or str(file_path), currency, fetched_at=retrieved)
        con.execute("INSERT INTO import_runs(file_path,sha256,symbol,snapshot_id,imported_at,invalid_rows,duplicate_rows) VALUES(?,?,?,?,?,?,?)", (str(file_path.resolve()), digest, symbol, identifier, datetime.now(timezone.utc).isoformat(), invalid, duplicates))
    return {"file": str(file_path), "symbol": symbol, "status": "imported", "snapshot_id": identifier, "rows": len(rows), "invalid_rows": invalid, "duplicate_rows": duplicates, "errors": samples, "first_date": rows[0]["date"], "last_date": rows[-1]["date"], "currency": currency}


def import_directory(path, directory):
    directory = Path(directory)
    if not directory.is_dir():
        raise ValueError("가격 원본 폴더가 없습니다.")
    reports = []
    # Direct source children only; generated workbooks in nested .work are not raw evidence.
    for file_path in sorted(directory.iterdir()):
        if not file_path.is_file() or file_path.suffix.lower() not in (".xlsx", ".csv"):
            continue
        try:
            reports.append(import_file(path, file_path))
        except (OSError, ValueError, TypeError, StopIteration) as exc:
            reports.append({"file": str(file_path), "status": "error", "error": str(exc)})
    return reports
